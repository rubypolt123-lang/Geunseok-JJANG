from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from investment_excel import sample
from investment_excel.models import Inputs, Ticker, Trade
from investment_excel.workbook import (
    CHART_DAYS,
    SAMPLE_NOTE,
    TICKER_FIRST,
    TRADE_FIRST,
    _nice_axis,
    build_workbook,
    read_workbook,
)

TODAY = date(2026, 10, 2)


def build(tmp_path, inputs=None, bars=None):
    inputs = inputs or sample.example_inputs(TODAY)
    bars = sample.example_bars(TODAY) if bars is None else bars
    path = tmp_path / "투자관리.xlsx"
    build_workbook(inputs, bars, updated_at="2026-10-02 09:00", data_note=SAMPLE_NOTE).save(path)
    return path


def test_sheets_and_order(tmp_path):
    wb = load_workbook(build(tmp_path))
    assert wb.sheetnames == ["사용법", "대시보드", "일봉차트", "보유현황", "거래기록", "종목", "설정", "일봉데이터", "차트데이터"]
    assert wb["차트데이터"].sheet_state == "hidden"
    assert wb.active.title == "대시보드"


def test_one_candle_and_volume_chart_per_ticker_with_data(tmp_path):
    bars = sample.example_bars(TODAY)
    del bars["AAPL"]
    wb = load_workbook(build(tmp_path, bars=bars))
    # openpyxl 은 읽을 때 차트를 버리므로 저장 전 객체로 확인
    built = build_workbook(sample.example_inputs(TODAY), bars)
    assert len(built["일봉차트"]._charts) == 2 * 3
    assert len(built["대시보드"]._charts) == 2
    links = [c.value for c in wb["일봉차트"][3] if c.value]
    assert links == ["▶ 비트코인", "▶ 이더리움", "▶ 삼성전자"]


def test_inputs_round_trip(tmp_path):
    inputs = Inputs(
        usd_krw=1385.5,
        auto_fx=False,
        tickers=[Ticker("005930", "삼성전자", "국내주식", "KRW"), Ticker("GOLD", "금", "기타", "KRW", manual_price=150000.0, memo="수동")],
        trades=[Trade(date(2026, 9, 1), "005930", "매수", 10, 70000, 100, "첫 매수"), Trade(date(2026, 8, 1), "GOLD", "매도", 1, 150000)],
    )
    bars = sample.example_bars(TODAY)
    path = build(tmp_path, inputs, {"005930": bars["005930"]})
    got, got_bars, warnings = read_workbook(path)
    assert warnings == []
    assert got.usd_krw == 1385.5 and got.auto_fx is False
    assert got.tickers == inputs.tickers
    assert sorted(got.trades, key=lambda t: t.day) == sorted(inputs.trades, key=lambda t: t.day)
    assert got_bars["005930"] == bars["005930"]


def test_leading_zero_codes_survive_number_entry(tmp_path):
    path = build(tmp_path)
    wb = load_workbook(path)
    wb["종목"][f"A{TICKER_FIRST + 4}"] = 660  # 사용자가 000660 을 숫자로 입력해 0 이 사라진 경우는 그대로 '660'
    wb["거래기록"][f"A{TRADE_FIRST + 10}"] = "2026-09-30"
    wb["거래기록"][f"B{TRADE_FIRST + 10}"] = "005930"
    wb["거래기록"][f"C{TRADE_FIRST + 10}"] = "매수"
    wb["거래기록"][f"D{TRADE_FIRST + 10}"] = 1
    wb["거래기록"][f"E{TRADE_FIRST + 10}"] = "61,000"
    wb.save(path)
    got, _, warnings = read_workbook(path)
    assert got.tickers[-1].code == "660"
    assert got.trades[-1] == Trade(date(2026, 9, 30), "005930", "매수", 1, 61000, 0.0, "")
    assert warnings == []


def test_bad_trade_rows_are_reported(tmp_path):
    path = build(tmp_path)
    wb = load_workbook(path)
    wb["거래기록"][f"B{TRADE_FIRST + 20}"] = "BTCUSDT"
    wb["거래기록"][f"C{TRADE_FIRST + 20}"] = "사기"
    wb.save(path)
    _, _, warnings = read_workbook(path)
    assert len(warnings) == 1 and f"{TRADE_FIRST + 20}행" in warnings[0]


def test_formulas_reference_expected_ranges(tmp_path):
    wb = load_workbook(build(tmp_path))
    h = wb["보유현황"]
    assert h["G5"].value == '=IF($A5="","",E5-F5)'
    assert "SUMIFS(거래기록!$D$5:$D$1004" in h["E5"].value
    assert h["R56"].value == "=SUM(R5:R54)"
    cd = wb["차트데이터"]
    assert cd["B1"].value == "BTCUSDT"
    assert cd["I125"].value.startswith("=IF(AND($B$2>0,$A125-19>=$B$2)")


def test_nice_axis_fits_window():
    bars = sample.example_bars(TODAY)["BTCUSDT"]
    low, high, step = _nice_axis(bars)
    window = bars[-CHART_DAYS:]
    assert low <= min(b.low for b in window) and high >= max(b.high for b in window)
    assert low > 0  # 0 부터 시작하지 않아 캔들이 납작해지지 않음
    assert (high - low) / step <= 10
