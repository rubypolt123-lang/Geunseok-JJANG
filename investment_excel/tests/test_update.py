from __future__ import annotations

from datetime import datetime

from openpyxl import load_workbook

from investment_excel import update
from investment_excel.prices import PriceError
from investment_excel.workbook import SAMPLE_NOTE, read_workbook

from .conftest import binance_payload, fake_get, naver_payload, yahoo_payload

NOW = datetime(2026, 10, 2, 9, 30)


def all_routes():
    return {
        "binance": binance_payload(300),
        "naver": naver_payload(200),
        "KRW%3DX": yahoo_payload(5, price=1380.0),
        "finance/chart/AAPL": yahoo_payload(200),
    }


def test_first_run_creates_file_with_real_prices(tmp_path, capsys):
    path = tmp_path / "투자관리.xlsx"
    assert update.run(path, get=fake_get(all_routes()), now=NOW) == 0
    inputs, bars, _ = read_workbook(path)
    assert [t.code for t in inputs.tickers] == ["BTCUSDT", "ETHUSDT", "005930", "AAPL"]
    assert len(bars["BTCUSDT"]) == 300 and len(bars["AAPL"]) == 199
    assert inputs.usd_krw == 1385.0
    settings = load_workbook(path)["설정"]
    assert settings["B6"].value == "2026-10-02 09:30"
    assert settings["B7"].value == update.REAL_NOTE
    assert not (tmp_path / update.BACKUP_DIR_NAME).exists()
    assert "✔ BTCUSDT" in capsys.readouterr().out


def test_offline_first_run_uses_sample(tmp_path):
    path = tmp_path / "투자관리.xlsx"
    assert update.run(path, offline=True, now=NOW) == 0
    assert load_workbook(path)["설정"]["B7"].value == SAMPLE_NOTE


def test_update_keeps_user_inputs_and_old_data_when_fetch_fails(tmp_path, capsys):
    path = tmp_path / "투자관리.xlsx"
    update.run(path, get=fake_get(all_routes()), now=NOW)
    wb = load_workbook(path)
    wb["거래기록"]["J5"] = "내 메모"
    wb["설정"]["B5"] = "아니오"
    wb["설정"]["B4"] = 1300
    wb.save(path)
    _, before, _ = read_workbook(path)

    routes = all_routes()
    routes["naver"] = PriceError("네이버 점검")
    assert update.run(path, get=fake_get(routes), now=NOW.replace(hour=10)) == 0

    inputs, bars, _ = read_workbook(path)
    assert inputs.trades[0].memo == "내 메모"
    assert inputs.usd_krw == 1300 and inputs.auto_fx is False  # 자동 환율 끔
    assert bars["005930"] == before["005930"]  # 받지 못한 종목은 기존 데이터 유지
    assert "005930: 네이버 점검 → 기존 데이터를 그대로 둡니다" in capsys.readouterr().out
    backups = list((tmp_path / update.BACKUP_DIR_NAME).glob("*.xlsx"))
    assert len(backups) == 1


def test_sample_note_stays_until_all_sample_data_is_replaced(tmp_path):
    path = tmp_path / "투자관리.xlsx"
    update.run(path, offline=True, now=NOW)
    routes = all_routes()
    routes["naver"] = PriceError("down")
    update.run(path, get=fake_get(routes), now=NOW.replace(hour=11))
    assert load_workbook(path)["설정"]["B7"].value == SAMPLE_NOTE  # 삼성전자는 아직 예시 데이터
    update.run(path, get=fake_get(all_routes()), now=NOW.replace(hour=12))
    assert load_workbook(path)["설정"]["B7"].value == update.REAL_NOTE


def test_locked_file_message(tmp_path, monkeypatch, capsys):
    path = tmp_path / "투자관리.xlsx"
    update.run(path, offline=True, now=NOW)

    def locked(src, dst):
        raise PermissionError("in use")

    monkeypatch.setattr(update.os, "replace", locked)
    assert update.run(path, offline=True, now=NOW.replace(hour=13)) == 1
    assert "엑셀에서 '투자관리.xlsx' 를 닫고" in capsys.readouterr().out
    assert not list(tmp_path.glob("~*"))


def test_backups_are_capped(tmp_path, monkeypatch):
    monkeypatch.setattr(update, "KEEP_BACKUPS", 2)
    path = tmp_path / "투자관리.xlsx"
    update.run(path, offline=True, now=NOW)
    for minute in range(4):
        update.run(path, offline=True, now=NOW.replace(minute=minute))
    assert len(list((tmp_path / update.BACKUP_DIR_NAME).glob("*.xlsx"))) == 2
